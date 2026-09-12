import os
import uuid
import openpyxl
from io import BytesIO
import datetime
from typing import List, Optional, Dict, Any
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Request
from fastapi.responses import Response, FileResponse
from sqlalchemy.orm import Session

from ..database import get_db
from ..models import Product, ProductVariant, SellerAccount, BuyboxTracking, PriceLog
from ..schemas import (
    ProductResponse,
    CurtainCalculateBatchRequest,
    CurtainCalculateItemResult,
    CreateCurtainProductRequestV2
)
from ..services.trendyol_client import (
    TrendyolClient, 
    TRENDYOL_PERDE_BRANDS, 
    TRENDYOL_PERDE_CATEGORIES, 
    TRENDYOL_CARGO_COMPANIES
)
from ..services.curtain_calculator import CurtainCalculator

router = APIRouter(prefix="/api/products", tags=["products"])

@router.get("", response_model=List[ProductResponse])
def get_products(db: Session = Depends(get_db)):
    products = db.query(Product).all()
    if not products:
        import asyncio
        asyncio.create_task(sync_products_from_source(db))
        products = db.query(Product).all()
    return products

@router.post("/upload-image")
async def upload_image(file: UploadFile = File(...), request: Request = None):
    """
    Kullanıcının bilgisayarından perde fotoğraflarını yükler ve URL döner.
    """
    allowed_extensions = {".jpg", ".jpeg", ".png", ".webp", ".gif"}
    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in allowed_extensions:
        raise HTTPException(status_code=400, detail="Desteklenmeyen dosya formatı. Lütfen JPG, PNG veya WEBP yükleyin.")

    uploads_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), "uploads")
    os.makedirs(uploads_dir, exist_ok=True)

    unique_filename = f"curtain_{uuid.uuid4().hex[:10]}{ext}"
    file_path = os.path.join(uploads_dir, unique_filename)

    with open(file_path, "wb") as f:
        content = await file.read()
        f.write(content)

    base_url = str(request.base_url) if request else "http://127.0.0.1:8000/"
    image_url = f"{base_url.rstrip('/')}/uploads/{unique_filename}"

    return {
        "status": "success",
        "filename": unique_filename,
        "url": image_url,
        "original_name": file.filename
    }

@router.post("/export-excel")
async def export_trendyol_excel(payload: CreateCurtainProductRequestV2):
    """
    Trendyol resmi 'tul-perde.xlsx' şablonuna birebir uygun 57 sütunlu Excel dosyası üretir.
    """
    template_path = r"C:\Users\murat\Downloads\tul-perde.xlsx"
    clean_model_code = payload.model_code.strip().upper().replace(" ", "-")

    if os.path.exists(template_path):
        wb = openpyxl.load_workbook(template_path)
    else:
        wb = openpyxl.Workbook()
        ws_default = wb.active
        ws_default.title = "Ürünlerinizi Burada Listeleyin"

    ws = wb["Ürünlerinizi Burada Listeleyin"] if "Ürünlerinizi Burada Listeleyin" in wb.sheetnames else wb.active

    # Row 1 is header, start inserting data from Row 2
    row_idx = 2
    for item in payload.variants:
        w = float(item.get("width_cm", 100))
        h = float(item.get("height_cm", 200))
        size_lbl = item.get("size_label", f"{int(w)} x {int(h)}")
        # Format for Trendyol Boyut/Ebat column (e.g. '100 x 260')
        boyut_ebat = f"{int(w)} x {int(h)}"
        barcode = item.get("barcode") or f"{clean_model_code}-{int(w)}-{int(h)}"
        sale_p = float(item.get("sale_price", 299.90))
        list_p = round(sale_p * 1.25, 2)
        stock_q = int(item.get("stock_quantity", 50))

        # 57 Columns mapping:
        ws.cell(row=row_idx, column=1, value=barcode) # Barkod
        ws.cell(row=row_idx, column=2, value=clean_model_code) # Model Kodu
        ws.cell(row=row_idx, column=3, value=payload.brand_name or "Taç") # Marka
        ws.cell(row=row_idx, column=4, value="895") # Kategori (Tül Perde: 895)
        ws.cell(row=row_idx, column=5, value="TRY") # Para Birimi
        ws.cell(row=row_idx, column=6, value=f"{payload.title} {size_lbl}") # Ürün Adı
        ws.cell(row=row_idx, column=7, value=payload.description or f"<p>{payload.title} birinci kalite dikişli perde.</p>") # Ürün Açıklaması
        ws.cell(row=row_idx, column=8, value=list_p) # Piyasa Satış Fiyatı
        ws.cell(row=row_idx, column=9, value=sale_p) # Trendyol Satış Fiyatı
        ws.cell(row=row_idx, column=10, value=stock_q) # Stok Adedi
        ws.cell(row=row_idx, column=11, value=barcode) # Stok Kodu
        ws.cell(row=row_idx, column=12, value=payload.vat_rate or 10) # KDV Oranı
        ws.cell(row=row_idx, column=13, value=0) # ÖTV Oranı
        ws.cell(row=row_idx, column=14, value=payload.dimensional_weight or 2) # Desi
        ws.cell(row=row_idx, column=16, value=payload.image_url or "") # Görsel 1
        ws.cell(row=row_idx, column=24, value=payload.delivery_duration or 3) # Sevkiyat Süresi
        ws.cell(row=row_idx, column=26, value=boyut_ebat) # Boyut/Ebat (Mavi Varyant Sütunu)
        ws.cell(row=row_idx, column=29, value="1") # Kanat Sayısı
        ws.cell(row=row_idx, column=42, value="Salon / Oturma Odası") # Kullanım Alanı
        ws.cell(row=row_idx, column=43, value=boyut_ebat) # Beden
        ws.cell(row=row_idx, column=44, value="Polyester") # Materyal
        ws.cell(row=row_idx, column=45, value="Normal (1 x 2.5)") # Pile
        ws.cell(row=row_idx, column=46, value="1") # Parça Sayısı
        ws.cell(row=row_idx, column=48, value="Kornişli") # Takma Şekli
        ws.cell(row=row_idx, column=49, value="Düz") # Desen
        ws.cell(row=row_idx, column=50, value="Şeffaf") # Işık Geçirgenliği
        ws.cell(row=row_idx, column=53, value="TR") # Menşei
        ws.cell(row=row_idx, column=54, value=payload.color or "Ekru") # Web Color
        ws.cell(row=row_idx, column=56, value=payload.color or "Ekru") # Renk
        
        row_idx += 1

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)

    filename = f"Trendyol_Yukleme_{clean_model_code}.xlsx"
    return Response(
        content=stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@router.get("/brands")
async def get_brands(name: Optional[str] = None, db: Session = Depends(get_db)):
    seller = db.query(SellerAccount).first()
    is_mock = seller.is_mock_mode if seller else True
    return await TrendyolClient.get_brands(name=name, is_mock=is_mock)

@router.get("/categories-tree")
async def get_categories(db: Session = Depends(get_db)):
    seller = db.query(SellerAccount).first()
    is_mock = seller.is_mock_mode if seller else True
    return await TrendyolClient.get_categories(is_mock=is_mock)

@router.get("/cargo-companies")
def get_cargo_companies():
    return TRENDYOL_CARGO_COMPANIES

@router.get("/categories/{category_id}/attributes")
async def get_category_attributes(category_id: int, db: Session = Depends(get_db)):
    seller = db.query(SellerAccount).first()
    is_mock = seller.is_mock_mode if seller else True
    return await TrendyolClient.get_category_attributes(category_id, is_mock=is_mock)

@router.get("/batch-status/{batch_request_id}")
async def get_batch_status(batch_request_id: str, db: Session = Depends(get_db)):
    seller = db.query(SellerAccount).first()
    return await TrendyolClient.get_batch_request_result(
        supplier_id=seller.supplier_id if seller else "DEMO",
        api_key=seller.api_key if seller else "DEMO",
        api_secret=seller.api_secret if seller else "DEMO",
        batch_request_id=batch_request_id,
        is_mock=seller.is_mock_mode if seller else True
    )

@router.post("/calculate-curtain", response_model=List[CurtainCalculateItemResult])
def calculate_curtain_sizes(payload: CurtainCalculateBatchRequest):
    results = []
    for s in payload.sizes:
        w = s.width
        h = s.height
        size_label = f"{int(w)} x {int(h)} cm"

        if payload.category_type == "tul":
            calc = CurtainCalculator.calculate_tul_perde(
                width_cm=w,
                height_cm=h,
                meter_price=payload.unit_price,
                pleat_type=payload.pleat_type or "1x2.5",
                cargo_cost=payload.cargo_cost or 45.0,
                commission_rate=payload.commission_rate or 0.20,
                target_profit_margin=payload.target_profit_margin or 0.25,
                round_to_90=payload.round_to_90 if payload.round_to_90 is not None else True
            )
            qty_or_m2 = calc["required_meters"]

        elif payload.category_type == "stor_zebra":
            calc = CurtainCalculator.calculate_stor_zebra(
                width_cm=w,
                height_cm=h,
                m2_price=payload.unit_price,
                has_skirt=payload.has_skirt or False,
                skirt_m2_extra=payload.skirt_m2_extra or 50.0,
                has_bead=payload.has_bead or False,
                bead_m2_extra=payload.bead_m2_extra or 50.0,
                cargo_cost=payload.cargo_cost or 45.0,
                commission_rate=payload.commission_rate or 0.20,
                target_profit_margin=payload.target_profit_margin or 0.25,
                round_to_90=payload.round_to_90 if payload.round_to_90 is not None else True
            )
            qty_or_m2 = calc["calculated_m2"]

        elif payload.category_type == "fon":
            calc = CurtainCalculator.calculate_fon_perde(
                width_cm=w,
                height_cm=h,
                meter_price=payload.unit_price,
                panel_type=payload.panel_type or "tek_kanat",
                pleat_type=payload.pleat_type or "1x2.5",
                cargo_cost=payload.cargo_cost or 45.0,
                commission_rate=payload.commission_rate or 0.20,
                target_profit_margin=payload.target_profit_margin or 0.25,
                round_to_90=payload.round_to_90 if payload.round_to_90 is not None else True
            )
            qty_or_m2 = calc["required_meters"]

        else: # karartma_saten
            calc = CurtainCalculator.calculate_blackout_saten(
                width_cm=w,
                height_cm=h,
                meter_price=payload.unit_price,
                cargo_cost=payload.cargo_cost or 45.0,
                commission_rate=payload.commission_rate or 0.20,
                target_profit_margin=payload.target_profit_margin or 0.25,
                round_to_90=payload.round_to_90 if payload.round_to_90 is not None else True
            )
            qty_or_m2 = calc["required_meters"]

        results.append(
            CurtainCalculateItemResult(
                width_cm=w,
                height_cm=h,
                size_label=size_label,
                calculated_quantity_or_m2=qty_or_m2,
                fabric_cost=calc["fabric_cost"],
                direct_cost=calc["direct_cost"],
                total_cost=calc["total_cost"],
                sale_price=calc["sale_price"],
                min_price=calc["min_price"],
                commission_amount=calc["commission_amount"],
                net_income=calc["net_income"],
                net_profit=calc["net_profit"],
                profit_margin_pct=calc["profit_margin_pct"]
            )
        )

    return results

@router.post("/create-v2")
async def create_product_v2(payload: CreateCurtainProductRequestV2, db: Session = Depends(get_db)):
    """
    Trendyol Ürün Yaratma V2 Standardında (POST /v2/products) ürün ve onlarca varyantı oluşturma.
    """
    seller = db.query(SellerAccount).first()
    if not seller:
        seller = SellerAccount(is_mock_mode=True)
        db.add(seller)
        db.commit()
        db.refresh(seller)

    clean_model_code = payload.model_code.strip().upper().replace(" ", "-")

    # 1. Yerel Veritabanına Kayıt
    prod = db.query(Product).filter(Product.model_code == clean_model_code).first()
    if not prod:
        prod = Product(
            seller_id=seller.id,
            model_code=clean_model_code,
            title=payload.title,
            brand=payload.brand_name,
            category_name=payload.category_name,
            image_url=payload.image_url or "https://images.unsplash.com/photo-1513694203232-719a280e022f?w=600&auto=format&fit=crop&q=80"
        )
        db.add(prod)
        db.commit()
        db.refresh(prod)
    else:
        prod.title = payload.title
        prod.brand = payload.brand_name
        prod.category_name = payload.category_name

    # 2. Trendyol V2 Payload Hazırlama
    v2_items = []

    for item in payload.variants:
        w = float(item.get("width_cm", 100))
        h = float(item.get("height_cm", 200))
        size_lbl = item.get("size_label", f"{int(w)} x {int(h)} cm")
        barcode = item.get("barcode") or f"{clean_model_code}-{int(w)}-{int(h)}"
        sale_p = float(item.get("sale_price", 299.90))
        list_p = round(sale_p * 1.25, 2)
        min_p = float(item.get("min_price", sale_p * 0.75))
        direct_cost = float(item.get("direct_cost", sale_p * 0.40))
        stock_q = int(item.get("stock_quantity", 50))

        # DB Varyant güncelle / ekle
        variant = db.query(ProductVariant).filter(ProductVariant.barcode == barcode).first()
        if not variant:
            variant = ProductVariant(
                product_id=prod.id,
                barcode=barcode,
                variant_name=size_lbl,
                width_cm=w,
                height_cm=h,
                fabric_type=payload.category_name,
                stock_quantity=stock_q,
                current_price=sale_p,
                cost_price=direct_cost,
                min_price=min_p,
                max_price=round(sale_p * 1.5, 2),
                cargo_cost=45.0,
                commission_rate=0.20
            )
            db.add(variant)
            db.commit()
            db.refresh(variant)

            buybox = BuyboxTracking(
                variant_id=variant.id,
                is_active=True,
                strategy="beat_by_diff",
                price_diff=0.50,
                has_buybox=True,
                winner_seller_name="Bizim Mağaza",
                last_buybox_price=sale_p
            )
            db.add(buybox)
        else:
            variant.current_price = sale_p
            variant.min_price = min_p
            variant.cost_price = direct_cost
            variant.stock_quantity = stock_q

        # Resim Listesi Hazırlama (Trendyol geçerli public URL ister)
        image_list = []
        raw_images = payload.images if (payload.images and len(payload.images) > 0) else ([payload.image_url] if payload.image_url else [])
        for u in raw_images:
            if u and isinstance(u, str) and u.strip():
                # Eğer localhost ise public fallback sağla
                if "127.0.0.1" in u or "localhost" in u:
                    image_list.append({"url": "https://images.unsplash.com/photo-1513694203232-719a280e022f?w=600&auto=format&fit=crop&q=80"})
                else:
                    image_list.append({"url": u.strip()})
        if not image_list:
            image_list = [{"url": "https://images.unsplash.com/photo-1513694203232-719a280e022f?w=600&auto=format&fit=crop&q=80"}]

        # Pile formatı standardizasyonu (Trendyol Excel değerleri)
        pleat_val = payload.pleat_type or "Normal (1 x 2.5)"
        if pleat_val == "1x2.5": pleat_val = "Normal (1 x 2.5)"
        elif pleat_val == "1x3": pleat_val = "Sık (1 x 3)"
        elif pleat_val == "1x2": pleat_val = "Seyrek (1 x 2)"
        elif pleat_val == "pilesiz": pleat_val = "Pilesiz (1 x 1)"

        # Kategoriye Göre Resmi Trendyol V2 Nitelik Listesi
        cat_id = payload.category_id or 895
        if cat_id == 1849: # Stor & Zebra Perde
            attributes_payload = [
                {"attributeId": 338, "customAttributeValue": f"{int(w)} x {int(h)}"}, # Beden
                {"attributeId": 34, "customAttributeValue": "Stor Perde"}, # Tip
                {"attributeId": 14, "customAttributeValue": payload.material or "Polyester"}, # Materyal
                {"attributeId": 33, "customAttributeValue": payload.pattern or "Düz"}, # Desen
                {"attributeId": 1192, "customAttributeValue": "TR"}, # Menşei
                {"attributeId": 348, "customAttributeValue": payload.color or "Ekru"}, # Web Color
                {"attributeId": 47, "customAttributeValue": payload.color or "Ekru"} # Renk
            ]
        else: # Tül Perde (895), Fon Perde (1848) ve Karartma
            attributes_payload = [
                {"attributeId": 92, "customAttributeValue": f"{int(w)} x {int(h)}"}, # Boyut/Ebat (Varyant Belirleyici)
                {"attributeId": 14, "customAttributeValue": payload.material or "Polyester"}, # Materyal
                {"attributeId": 1101, "customAttributeValue": pleat_val}, # Pile
                {"attributeId": 18, "customAttributeValue": "1"}, # Parça Sayısı
                {"attributeId": 258, "customAttributeValue": payload.hanging_type or "Kornişli"}, # Takma Şekli
                {"attributeId": 33, "customAttributeValue": payload.pattern or "Düz"}, # Desen
                {"attributeId": 1192, "customAttributeValue": "TR"}, # Menşei
                {"attributeId": 348, "customAttributeValue": payload.color or "Ekru"}, # Web Color
                {"attributeId": 47, "customAttributeValue": payload.color or "Ekru"} # Renk
            ]

        # Resmi Trendyol V2 Item Şeması
        v2_item = {
            "barcode": barcode,
            "title": f"{payload.title} {size_lbl}",
            "productMainId": clean_model_code,
            "brandId": payload.brand_id,
            "categoryId": cat_id,
            "quantity": stock_q,
            "stockCode": barcode,
            "dimensionalWeight": payload.dimensional_weight or 2.0,
            "description": payload.description or f"<p>{payload.title} kaliteli kumaş perde.</p>",
            "currencyType": "TRY",
            "listPrice": list_p,
            "salePrice": sale_p,
            "vatRate": payload.vat_rate or 10,
            "cargoCompanyId": payload.cargo_company_id or 10,
            "deliveryDuration": payload.delivery_duration or 2,
            "images": image_list,
            "attributes": attributes_payload
        }
        v2_items.append(v2_item)

    db.commit()

    # 3. Trendyol V2 Endpoint'ine Gönder (POST /integration/product/sellers/{sellerId}/v2/products)
    api_resp = await TrendyolClient.create_products_v2(
        supplier_id=seller.supplier_id,
        api_key=seller.api_key,
        api_secret=seller.api_secret,
        items=v2_items,
        is_mock=seller.is_mock_mode
    )

    return {
        "status": "success",
        "product_id": prod.id,
        "model_code": clean_model_code,
        "variant_count": len(v2_items),
        "batch_request_id": api_resp.get("batchRequestId"),
        "message": f"Trendyol V2: '{prod.title}' ürünü {len(v2_items)} adet ölçü varyantıyla başarıyla Trendyol API V2'ye gönderildi.",
        "trendyol_response": api_resp
    }

async def sync_products_from_source(db: Session):
    seller = db.query(SellerAccount).first()
    if not seller:
        seller = SellerAccount(is_mock_mode=True)
        db.add(seller)
        db.commit()
        db.refresh(seller)

    raw_products = await TrendyolClient.fetch_products(
        supplier_id=seller.supplier_id,
        api_key=seller.api_key,
        api_secret=seller.api_secret,
        is_mock=seller.is_mock_mode
    )

    synced_count = 0
    variant_count = 0

    for item in raw_products:
        model_code = item.get("model_code") or item.get("productCode", "DEMO-01")
        prod = db.query(Product).filter(Product.model_code == model_code).first()
        if not prod:
            prod = Product(
                seller_id=seller.id,
                model_code=model_code,
                title=item.get("title", "Örnek Ürün"),
                brand=item.get("brand", "Taç"),
                category_name=item.get("category_name", "Perde"),
                image_url=item.get("image_url")
            )
            db.add(prod)
            db.commit()
            db.refresh(prod)
            synced_count += 1

        for var_data in item.get("variants", []):
            barcode = var_data.get("barcode")
            variant = db.query(ProductVariant).filter(ProductVariant.barcode == barcode).first()
            if not variant:
                variant = ProductVariant(
                    product_id=prod.id,
                    barcode=barcode,
                    variant_name=var_data.get("name", "Standart"),
                    width_cm=float(var_data.get("w", 140)),
                    height_cm=float(var_data.get("h", 260)),
                    current_price=float(var_data.get("price", 299.90)),
                    cost_price=float(var_data.get("cost", 120.0)),
                    min_price=float(var_data.get("min_price", 219.90)),
                    max_price=float(var_data.get("max_price", 450.0)),
                    stock_quantity=int(var_data.get("stock", 50))
                )
                db.add(variant)
                db.commit()
                db.refresh(variant)
                variant_count += 1

                buybox = BuyboxTracking(
                    variant_id=variant.id,
                    is_active=True,
                    strategy="beat_by_diff",
                    price_diff=0.50,
                    has_buybox=True,
                    winner_seller_name="Bizim Mağaza",
                    last_buybox_price=variant.current_price
                )
                db.add(buybox)
                db.commit()

    return {
        "status": "success",
        "message": f"{synced_count} ana ürün ve {variant_count} varyant eşitlendi.",
        "is_mock": seller.is_mock_mode
    }
